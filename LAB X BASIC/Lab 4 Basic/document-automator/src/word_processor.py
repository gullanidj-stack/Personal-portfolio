import os

import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .utils import setup_logging

logger = setup_logging(__name__)

_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


class WordProcessor:
    """Handles Word (.docx) operations using python-docx."""

    def __init__(self):
        logger.info("WordProcessor initialized.")

    def create_document(self):
        document = docx.Document()
        logger.info("New Word document created.")
        return document

    def add_heading(self, document, text, level=1):
        document.add_heading(text, level=level)
        logger.info(f"Added heading (Level {level}): '{text}'")

    def add_paragraph(self, document, text, style=None, bold=False,
                      italic=False, font_size=None, align="left"):
        paragraph = document.add_paragraph(text, style=style)
        for run in paragraph.runs:
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if font_size:
                run.font.size = Pt(font_size)
        paragraph.alignment = _ALIGNMENTS.get(align, WD_ALIGN_PARAGRAPH.LEFT)
        logger.info(f"Added paragraph: '{text[:50]}...'")

    def add_table(self, document, data, include_header=True,
                  header_bold=True, col_widths=None):
        """`data` is a list of rows; the first row is the header if include_header."""
        if not data or not data[0]:
            logger.warning("Table data is empty or malformed. Skipping table creation.")
            return

        rows, cols = len(data), len(data[0])
        table = document.add_table(rows=rows, cols=cols)
        table.style = "Table Grid"

        for r_idx, row_data in enumerate(data):
            for c_idx, cell_data in enumerate(row_data):
                cell = table.cell(r_idx, c_idx)
                cell.text = str(cell_data)
                if include_header and header_bold and r_idx == 0:
                    for run in cell.paragraphs[0].runs:
                        run.bold = True

        if col_widths and len(col_widths) == cols:
            for c_idx, width in enumerate(col_widths):
                for cell in table.columns[c_idx].cells:
                    cell.width = Inches(width)

        logger.info(f"Added table with {rows} rows and {cols} columns.")

    def add_picture(self, document, image_path, width_inches=None):
        if not os.path.exists(image_path):
            logger.error(f"Image file not found: {image_path}")
            return False
        try:
            if width_inches:
                document.add_picture(image_path, width=Inches(width_inches))
            else:
                document.add_picture(image_path)
            logger.info(f"Added picture: {os.path.basename(image_path)}")
            return True
        except Exception as e:
            logger.error(f"Error adding picture '{image_path}': {e}")
            return False

    def save_document(self, document, output_path):
        try:
            document.save(output_path)
            logger.info(f"Word document saved to '{output_path}'.")
            return True
        except Exception as e:
            logger.error(f"Error saving Word document to '{output_path}': {e}")
            return False
