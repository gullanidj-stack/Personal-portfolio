"""สร้างไฟล์ตั้งต้น: templates/report_template.docx และ PDF ตัวอย่างใน documents/
รัน:  python scripts/make_samples.py
"""
import sys
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

BASE = Path(__file__).resolve().parent.parent
FONT = "Tahoma"  # รองรับภาษาไทย และมีในทั้ง Windows / macOS


def style_run(run, size=12, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)


def add(document, text, size=12, bold=False, align=None):
    p = document.add_paragraph()
    style_run(p.add_run(text), size, bold)
    if align:
        p.alignment = align
    return p


def make_template(path):
    d = docx.Document()
    add(d, "{{REPORT_TITLE}}", 20, True, WD_ALIGN_PARAGRAPH.CENTER)
    add(d, "วันที่: {{REPORT_DATE}}")
    add(d, "ผู้จัดทำ: {{AUTHOR_NAME}}")
    add(d, "หน่วยงาน: {{COMPANY_NAME}}")
    add(d, "─" * 40)
    add(d, "บทนำ", 14, True)
    add(d, "{{INTRODUCTION_TEXT}}")
    add(d, "ภาพรวมประจำเดือน: {{MONTH_PERIOD}}", 14, True)
    add(d, "ตารางสรุปตัวชี้วัดสำคัญ:")
    add(d, "{{TABLE_PLACEHOLDER}}")
    add(d, "สรุป", 14, True)
    add(d, "รายงานฉบับนี้สิ้นสุดการวิเคราะห์ผลการดำเนินงาน")
    footer = d.sections[0].footer.paragraphs[0]
    style_run(footer.add_run("{{COMPANY_NAME}} | สร้างอัตโนมัติด้วย Document Agent"), 9)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(path)


def make_pdf(path, title, pages):
    path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    for n, body in enumerate(pages, start=1):
        c.setFont("Helvetica-Bold", 18)
        c.drawString(60, h - 80, title)
        c.setFont("Helvetica", 12)
        y = h - 120
        for line in body:
            c.drawString(60, y, line)
            y -= 20
        c.setFont("Helvetica", 9)
        c.drawString(60, 40, f"Page {n} of {len(pages)}")
        c.showPage()
    c.save()


def main():
    make_template(BASE / "templates" / "report_template.docx")
    docs = BASE / "documents"
    make_pdf(docs / "source_document.pdf", "Source Document", [
        ["This is the source document used for watermarking and text extraction.",
         "Page one introduces the project goals and scope."],
        ["Page two lists the deliverables and the project timeline."],
    ])
    make_pdf(docs / "dummy_doc1.pdf", "Project Documentation - Part A", [
        ["Architecture overview and design principles."],
    ])
    make_pdf(docs / "dummy_doc2.pdf", "Project Documentation - Part B", [
        ["Usage guide: how to write a JSON config."],
        ["Appendix: troubleshooting and FAQ."],
    ])
    print("สร้างไฟล์ตัวอย่างเรียบร้อย")


if __name__ == "__main__":
    sys.exit(main())
