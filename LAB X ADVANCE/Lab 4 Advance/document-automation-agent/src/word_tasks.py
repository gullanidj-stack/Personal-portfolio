"""งานเกี่ยวกับไฟล์ Word: เติม template, สร้างตาราง, ดึงข้อความ"""
import csv
import re
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .utils import ensure_directory_exists, setup_logging

logger = setup_logging(__name__)

TABLE_ANCHOR = "{{TABLE_PLACEHOLDER}}"
PLACEHOLDER_RE = re.compile(r"\{\{[A-Z0-9_]+\}\}")


class WordTasks:
    # ---------- helpers ----------
    @staticmethod
    def _iter_paragraphs(document):
        """วนทุก paragraph: เนื้อหา, ในตาราง, header และ footer"""
        yield from document.paragraphs
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs
        for section in document.sections:
            for part in (section.header, section.footer):
                yield from part.paragraphs

    @staticmethod
    def _fill_paragraph(paragraph, mapping):
        original = paragraph.text
        if "{{" not in original:
            return
        updated = original
        for key, value in mapping.items():
            updated = updated.replace("{{%s}}" % key, str(value))
        if updated == original:
            return
        # ข้อความถูกแบ่งเป็นหลาย run ได้ -> รวมไว้ใน run แรกเพื่อไม่ให้ placeholder ขาดตอน
        runs = paragraph.runs
        if runs:
            runs[0].text = updated
            for run in runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(updated)

    @staticmethod
    def _shade(cell, hex_fill):
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_fill)
        tc_pr.append(shd)

    def _build_table(self, document, rows):
        n_cols = max(len(r) for r in rows)
        table = document.add_table(rows=0, cols=n_cols)
        try:
            table.style = "Table Grid"
        except (KeyError, ValueError):
            logger.warning("template ไม่มี style 'Table Grid' -> ใช้ style เริ่มต้น")
        for i, row_data in enumerate(rows):
            cells = table.add_row().cells
            for j in range(n_cols):
                cell = cells[j]
                cell.text = str(row_data[j]) if j < len(row_data) else ""
                if i == 0:  # แถวหัวตาราง
                    for run in cell.paragraphs[0].runs:
                        run.bold = True
                    self._shade(cell, "D9E2F3")
        return table

    @staticmethod
    def _read_csv(path):
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [row for row in csv.reader(f) if row]

    # ---------- public tasks ----------
    def generate_report_from_template(self, template_path, output_path, placeholders,
                                      table_data=None, table_csv=None):
        if not Path(template_path).exists():
            logger.error("ไม่พบ template: %s", template_path)
            return False
        if table_csv:
            if not Path(table_csv).exists():
                logger.error("ไม่พบไฟล์ CSV ของตาราง: %s", table_csv)
                return False
            table_data = self._read_csv(table_csv)

        document = docx.Document(template_path)

        for paragraph in self._iter_paragraphs(document):
            self._fill_paragraph(paragraph, placeholders)

        anchor = next((p for p in document.paragraphs if TABLE_ANCHOR in p.text), None)
        if table_data:
            table = self._build_table(document, table_data)
            if anchor is not None:
                # ย้ายตารางไปไว้ตรงตำแหน่ง placeholder (ตัวอย่างเดิมตารางไปต่อท้ายเอกสาร)
                anchor._p.addnext(table._tbl)
            else:
                logger.warning("ไม่พบ %s ใน template -> ตารางถูกต่อท้ายเอกสาร", TABLE_ANCHOR)
        if anchor is not None:
            anchor._p.getparent().remove(anchor._p)

        leftover = {m for p in self._iter_paragraphs(document)
                    for m in PLACEHOLDER_RE.findall(p.text)}
        if leftover:
            logger.warning("ยังมี placeholder ที่ไม่ได้ถูกแทนที่: %s", ", ".join(sorted(leftover)))

        ensure_directory_exists(Path(output_path).parent)
        document.save(output_path)
        logger.info("สร้างรายงานแล้ว: %s", output_path)
        return True

    def extract_text_from_word(self, docx_path):
        """ดึงข้อความตามลำดับจริงในเอกสาร (ย่อหน้าและตารางสลับกันถูกตำแหน่ง)"""
        if not Path(docx_path).exists():
            logger.error("ไม่พบไฟล์ Word: %s", docx_path)
            return None
        document = docx.Document(docx_path)
        lines = []
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                lines.append(Paragraph(child, document).text)
            elif child.tag == qn("w:tbl"):
                for row in Table(child, document).rows:
                    lines.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(lines)
