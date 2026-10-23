# src/spreadsheet_agent.py
import json
import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, IconSetRule, FormulaRule
from openpyxl.utils import get_column_letter, range_boundaries
from .utils import setup_logging

logger = setup_logging(__name__)


class SpreadsheetAgent:
    """Performs a series of operations on an Excel workbook based on a configuration."""

    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.audit_log = []
        self.start_time = datetime.datetime.now()
        logger.info("Spreadsheet Agent initialized.")

    # ---------- helpers ----------
    def _log_audit(self, status, message, task_type="N/A", details=None):
        self.audit_log.append({
            "timestamp": datetime.datetime.now().isoformat(),
            "status": status,
            "task_type": task_type,
            "message": message,
            "details": details if details else {}
        })
        if status in ("ERROR", "CRITICAL"):
            logger.error(f"AUDIT LOG - {status}: {message} | Task: {task_type} | Details: {details}")
        else:
            logger.info(f"AUDIT LOG - {status}: {message} | Task: {task_type}")

    def _load_workbook(self):
        input_file = self.config["input_file"]
        try:
            self.workbook = openpyxl.load_workbook(input_file)
            self._log_audit("SUCCESS", f"Workbook '{input_file}' loaded.", "load_workbook")
            return True
        except FileNotFoundError:
            self._log_audit("ERROR", f"Input file not found: {input_file}", "load_workbook")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to load workbook '{input_file}': {e}", "load_workbook")
        return False

    def _get_sheet(self, sheet_name):
        """Gets a worksheet, creating it if it does not exist."""
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        self._log_audit("WARNING", f"Sheet '{sheet_name}' not found. Creating it.", "get_sheet")
        return self.workbook.create_sheet(sheet_name)

    # ---------- tasks ----------
    def _process_copy_data_task(self, task):
        src_name, dest_name = task.get("source_sheet"), task.get("dest_sheet")
        if not src_name or not dest_name:
            self._log_audit("ERROR", "Missing source_sheet or dest_sheet.", "copy_data", task)
            return
        if src_name not in self.workbook.sheetnames:
            self._log_audit("ERROR", f"Source sheet '{src_name}' not found.", "copy_data", task)
            return
        src, dest = self.workbook[src_name], self._get_sheet(dest_name)
        for row in src.iter_rows():
            for cell in row:
                dest[cell.coordinate].value = cell.value
        for letter, dim in src.column_dimensions.items():
            dest.column_dimensions[letter].width = dim.width
        self._log_audit("SUCCESS", f"Data copied from '{src_name}' to '{dest_name}'.", "copy_data")

    def _process_calculate_column_task(self, task):
        sheet_name = task.get("sheet")
        col = task.get("target_column")
        start_row = task.get("start_row", 2)
        formula = task.get("formula")
        if not sheet_name or not col or not formula:
            self._log_audit("ERROR", "Missing required fields for calculate_column.", "calculate_column", task)
            return
        sheet = self._get_sheet(sheet_name)
        sheet[f"{col}1"].value = task.get("header", "Calculated Column")
        sheet[f"{col}1"].font = Font(bold=True)
        for r in range(start_row, sheet.max_row + 1):
            try:
                cell = sheet[f"{col}{r}"]
                cell.value = formula.replace("{row}", str(r))
                if task.get("number_format"):
                    cell.number_format = task["number_format"]
            except Exception as e:
                self._log_audit("ERROR", f"Failed to apply formula to {col}{r}: {e}", "calculate_column", task)
        sheet.column_dimensions[col].width = max(14, len(task.get("header", "")) + 4)
        self._log_audit("SUCCESS", f"Column '{col}' calculated in '{sheet_name}'.", "calculate_column")

    def _process_conditional_format_task(self, task):
        sheet_name, rng, rule_type = task.get("sheet"), task.get("range"), task.get("rule_type")
        if not sheet_name or not rng or not rule_type:
            self._log_audit("ERROR", "Missing required fields for conditional_format.", "conditional_format", task)
            return
        sheet = self._get_sheet(sheet_name)
        try:
            if rule_type == "color_scale":
                rule = ColorScaleRule(
                    start_type='min', start_color=task.get("min_color", "FF0000"),
                    mid_type='percentile', mid_value=50, mid_color=task.get("mid_color", "FFFF00"),
                    end_type='max', end_color=task.get("max_color", "00FF00"))
            elif rule_type == "data_bar":
                rule = DataBarRule(start_type='min', end_type='max', color=task.get("color", "0070C0"))
            elif rule_type == "icon_set":
                rule = IconSetRule(task.get("icon_set", "3TrafficLights1"),
                                   task.get("value_type", "percent"),
                                   task.get("values", [0, 33, 67]))
            elif rule_type == "expression":
                formula = task.get("formula", "").lstrip("=")  # CF formulas have no leading '='
                fill_color = task.get("fill_color", "FFC7CE")
                rule = FormulaRule(
                    formula=[formula],
                    font=Font(color=task.get("font_color", "9C0006")),
                    fill=PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid"))
            else:
                self._log_audit("ERROR", f"Unsupported rule_type: {rule_type}", "conditional_format", task)
                return
            sheet.conditional_formatting.add(rng, rule)
            self._log_audit("SUCCESS", f"Conditional formatting '{rule_type}' applied to '{rng}' in '{sheet_name}'.",
                            "conditional_format")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to apply conditional formatting to '{rng}': {e}",
                            "conditional_format", task)

    def _process_chart_task(self, task):
        sheet_name, chart_type, data_range = task.get("sheet"), task.get("chart_type"), task.get("data_range")
        if not sheet_name or not chart_type or not data_range:
            self._log_audit("ERROR", "Missing required fields for chart task.", "chart", task)
            return
        sheet = self._get_sheet(sheet_name)
        title = task.get("title", "Chart Title")
        try:
            charts = {"bar": BarChart, "line": LineChart, "pie": PieChart}
            if chart_type not in charts:
                self._log_audit("ERROR", f"Unsupported chart type: {chart_type}", "chart", task)
                return
            chart = charts[chart_type]()

            c1, r1, c2, r2 = range_boundaries(data_range)
            data = Reference(sheet, min_col=c1, min_row=r1, max_col=c2, max_row=r2)
            chart.add_data(data, titles_from_data=True)  # first row of range = series title

            if task.get("category_range"):
                cc1, cr1, cc2, cr2 = range_boundaries(task["category_range"])
                chart.set_categories(Reference(sheet, min_col=cc1, min_row=cr1, max_col=cc2, max_row=cr2))

            chart.title = title
            if chart_type != "pie":
                chart.x_axis.title = task.get("x_axis_title", "")
                chart.y_axis.title = task.get("y_axis_title", "")
                # openpyxl >= 3.1 hides axes unless delete is False
                chart.x_axis.delete = False
                chart.y_axis.delete = False
            chart.width, chart.height = 18, 9

            sheet.add_chart(chart, task.get("top_left_cell", "G2"))
            self._log_audit("SUCCESS", f"Chart '{title}' ({chart_type}) added to '{sheet_name}'.", "chart")
        except Exception as e:
            self._log_audit("ERROR", f"Failed to create chart '{title}': {e}", "chart", task)

    def _add_audit_log_sheet(self):
        if not self.audit_log:
            return
        ws = self.workbook.create_sheet("Audit Log", 0)
        ws.append(["Timestamp", "Status", "Task Type", "Message", "Details"])
        for c in range(1, 6):
            cell = ws.cell(row=1, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="333333", end_color="333333", fill_type="solid")
        for e in self.audit_log:
            ws.append([e["timestamp"], e["status"], e["task_type"], e["message"],
                       json.dumps(e["details"], ensure_ascii=False)])
        for c in range(1, 6):
            letter = get_column_letter(c)
            longest = max((len(str(cell.value)) for cell in ws[letter] if cell.value), default=0)
            ws.column_dimensions[letter].width = min(longest + 2, 80)
        logger.info("Audit log sheet created.")

    # ---------- run ----------
    def run(self):
        if not self._load_workbook():
            self._log_audit("CRITICAL", "Agent could not start: workbook loading failed.", "run")
            return
        try:
            handlers = {
                "copy_data": self._process_copy_data_task,
                "calculate_column": self._process_calculate_column_task,
                "conditional_format": self._process_conditional_format_task,
                "chart": self._process_chart_task,
            }
            for i, task in enumerate(self.config["tasks"], start=1):
                task_type = task.get("type")
                logger.info(f"Executing Task {i}: Type='{task_type}' Sheet='{task.get('sheet')}'")
                handler = handlers.get(task_type)
                if handler:
                    handler(task)
                else:
                    self._log_audit("WARNING", f"Unsupported task type: {task_type}", "task_execution", task)

            self._log_audit("INFO", "All tasks processed.", "run")
            self._add_audit_log_sheet()

            output_file = self.config["output_file"]
            try:
                self.workbook.save(output_file)
                logger.info(f"Final workbook saved to: {output_file}")
            except Exception as e:
                logger.critical(f"Error saving output workbook (is it open in Excel?): {e}")
        except Exception as e:
            self._log_audit("CRITICAL", f"Unhandled error during agent execution: {e}", "run")
            logger.critical(f"An unhandled error occurred: {e}", exc_info=True)
        finally:
            duration = (datetime.datetime.now() - self.start_time).total_seconds()
            logger.info(f"Agent finished. Total time: {duration:.2f} seconds.")
